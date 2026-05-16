from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Iterable

from app.models import DonationRecord
from app.parser import parse_raw_pages, record_from_mapping

FIELD_ALIASES = {
    "record_id": ("记录号", "id", "record_id"),
    "donate_time": ("捐赠时间", "donateTime", "donate_time"),
    "donor": ("捐赠人", "donor"),
    "major": ("院系专业", "major"),
    "alumni_assoc": ("校友会", "alumniAssoc", "alumni_assoc"),
    "project_name": ("捐赠项目", "projName", "project_name"),
    "donate_amount": ("捐赠金额", "donateAmt", "donate_amount"),
}


def load_local_records(path: str | Path) -> list[DonationRecord]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".json":
        return _load_json_records(source)
    if suffix == ".csv":
        return _records_from_rows(_read_csv_rows(source))
    if suffix == ".xlsx":
        return _records_from_rows(_read_xlsx_rows(source))
    raise ValueError(f"Unsupported local data file: {source}")


def maybe_load_existing_records(raw_path: str | Path, db_path: str | Path) -> list[DonationRecord]:
    raw_source = Path(raw_path)
    if raw_source.exists():
        return load_local_records(raw_source)

    database_source = Path(db_path)
    if database_source.exists():
        from app.database import DonationDatabase

        return DonationDatabase(database_source).fetch_all_records()
    return []


def _load_json_records(path: Path) -> list[DonationRecord]:
    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)

    if _is_raw_pages(data):
        return parse_raw_pages(data)
    rows = data["records"] if isinstance(data, dict) and isinstance(data.get("records"), list) else data
    return _records_from_rows(rows)


def _is_raw_pages(data: Any) -> bool:
    return isinstance(data, list) and any(
        isinstance(page, dict) and isinstance(page.get("result"), dict) for page in data
    )


def _read_csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def _read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            return []
        header_names = ["" if header is None else str(header).strip() for header in headers]
        return [
            dict(zip(header_names, row))
            for row in rows
            if any(value is not None and value != "" for value in row)
        ]
    finally:
        workbook.close()


def _records_from_rows(rows: Iterable[dict[str, Any]]) -> list[DonationRecord]:
    return [record_from_mapping(_normalize_row(row)) for row in rows]


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for field, aliases in FIELD_ALIASES.items():
        normalized[field] = next((row[alias] for alias in aliases if alias in row), None)
    return normalized
